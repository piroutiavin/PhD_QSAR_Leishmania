"""
Manuscript regeneration (Word .docx) from the honest rebuild results
====================================================================
Generates JCIM_manuscript_rebuilt.docx from the leakage-free per-organism
sulfonamide panel. Numbers are pulled from results/ where practical; figures are
embedded from figures/sulfonamide_panel/. Virtual screening is intentionally
deferred to future work (not part of the rebuild).

Run:  .venv/Scripts/python.exe -m src.build_manuscript
"""

from __future__ import annotations
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPORT = Path("results/report/master_results.csv")
SAR = Path("results/shap_panel/L_infantum_shap_importance.csv")
FIG = Path("figures/sulfonamide_panel")
OUT = Path("JCIM_manuscript_rebuilt.docx")


def body(doc, text, size=11, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         space=6, bold=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(size)
    r.italic = italic; r.bold = bold
    return p


def h1(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    r = p.add_run(text.upper()); r.bold = True
    r.font.name = "Times New Roman"; r.font.size = Pt(12)


def h2(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run(text); r.bold = True
    r.font.name = "Times New Roman"; r.font.size = Pt(11)


def figure(doc, path, caption, width=5.8):
    if Path(path).exists():
        doc.add_picture(str(path), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        c.paragraph_format.space_after = Pt(10)
        r = c.add_run(caption); r.font.size = Pt(9); r.font.name = "Times New Roman"


def results_table(doc, mt: pd.DataFrame):
    cols = ["organism", "n", "active_pct", "final_model", "MCC", "ROC_AUC",
            "AD_inside_%", "Yrand_p"]
    headers = ["Organism", "n", "Active %", "Final model", "MCC", "ROC-AUC",
               "AD %", "Y-rand p"]
    t = doc.add_table(rows=1, cols=len(cols)); t.style = "Light Grid Accent 1"
    for i, hd in enumerate(headers):
        cell = t.rows[0].cells[i]; cell.text = ""
        rr = cell.paragraphs[0].add_run(hd); rr.bold = True; rr.font.size = Pt(9)
    for _, row in mt.iterrows():
        cells = t.add_row().cells
        for i, ckey in enumerate(cols):
            cells[i].text = str(row[ckey])
            for pp in cells[i].paragraphs:
                for rr in pp.runs:
                    rr.font.size = Pt(9)


def build():
    mt = pd.read_csv(REPORT)
    sar = pd.read_csv(SAR)
    top_sar = ", ".join(sar["descriptor"].head(6))
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)

    # ── Title / authors ──
    tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run("A Leakage-Free Machine-Learning QSAR Panel for "
                    "Anti-Kinetoplastid Sulfonamides, with Leishmania infantum "
                    "as the Primary Target")
    tr.bold = True; tr.font.size = Pt(14); tr.font.name = "Times New Roman"
    body(doc, "Avin", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space=2, bold=True)
    body(doc, "Affiliation to be added", size=10, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space=10)

    # ── Abstract ──
    h1(doc, "Abstract")
    linf = mt[mt.organism.str.contains("infantum")].iloc[0]
    body(doc,
        "Leishmaniasis and Chagas disease are neglected tropical diseases caused "
        "by kinetoplastid parasites for which safer chemotherapies are urgently "
        "needed. Sulfonamides are a synthetically tractable scaffold of interest, "
        "but public bioactivity data are scarce and heterogeneous, making honest "
        "model building difficult. We report a rigorously leakage-free quantitative "
        "structure-activity relationship (QSAR) panel for sulfonamide derivatives "
        "against four kinetoplastid parasites, with Leishmania infantum as the "
        "primary target. Sulfonamide compounds were extracted from ChEMBL and "
        "curated into four per-organism datasets (L. infantum, n=142; L. donovani, "
        "n=238; Trypanosoma cruzi, n=398; L. amazonensis, n=72). To avoid the "
        "optimistic bias that arises when the same molecule appears in both the "
        "training and test partitions, all models were evaluated by repeated, "
        "scaffold-grouped, activity-stratified cross-validation with feature "
        "selection refit inside every fold. A full algorithm panel (logistic "
        "regression, SVM, random forest, gradient boosting, XGBoost, LightGBM, and "
        "consensus/stacking ensembles) was compared; no ensemble or gradient-"
        "boosting method outperformed the best single interpretable model within "
        f"cross-validation uncertainty. The primary L. infantum model (logistic "
        f"regression on 12 interpretable Mordred descriptors) achieved a Matthews "
        f"correlation coefficient of {linf['MCC']} and ROC-AUC of {linf['ROC_AUC']}, "
        "with 98.6% applicability-domain coverage and significant Y-randomization "
        "(p=0.02). SHAP analysis revealed largely organism-specific structure-"
        "activity relationships. We present these modest but honest results as a "
        "reproducible baseline for anti-kinetoplastid sulfonamide discovery.",
        size=10)
    body(doc, "Keywords: QSAR; sulfonamides; Leishmania infantum; kinetoplastids; "
              "scaffold split; data leakage; SHAP; neglected tropical diseases.",
         size=10, italic=True)

    # ── Introduction ──
    h1(doc, "1. Introduction")
    h2(doc, "1.1 Kinetoplastid neglected tropical diseases")
    body(doc,
        "The kinetoplastid parasites Leishmania spp. and Trypanosoma cruzi cause "
        "leishmaniasis and Chagas disease, together affecting millions of people "
        "and disproportionately burdening low-income regions. Visceral "
        "leishmaniasis, caused chiefly by L. infantum and L. donovani, is fatal if "
        "untreated. Current drugs suffer from toxicity, resistance, and difficult "
        "administration, motivating the search for new chemical starting points.")
    h2(doc, "1.2 Sulfonamides and computational discovery")
    body(doc,
        "Sulfonamides are a versatile, synthetically accessible pharmacophore with "
        "documented antiparasitic activity. Quantitative structure-activity "
        "relationship (QSAR) modeling offers a low-cost route to prioritize such "
        "compounds. However, QSAR built on aggregated public data is easily "
        "over-optimistic: when a random train/test split allows the same compound "
        "(or near-identical scaffolds) to appear on both sides, reported metrics no "
        "longer reflect prospective performance.")
    h2(doc, "1.3 Objectives")
    body(doc,
        "We build and honestly validate a per-organism QSAR panel for sulfonamides "
        "against four kinetoplastids, treating L. infantum as the primary target. "
        "Our aims are to (i) curate sulfonamide-specific, per-organism datasets; "
        "(ii) evaluate a full algorithm panel under a strictly leakage-free "
        "scaffold-grouped protocol; (iii) select parsimonious, interpretable final "
        "models; and (iv) derive SHAP-based structure-activity relationships. We "
        "deliberately report modest, reproducible metrics rather than the inflated "
        "values obtainable from leaky splits.")

    # ── Methods ──
    h1(doc, "2. Materials and Methods")
    h2(doc, "2.1 Data collection and sulfonamide filtering")
    body(doc,
        "Bioactivity records (IC50/EC50, nM, relation '=') for L. infantum "
        "(CHEMBL612848), L. donovani (CHEMBL367), L. amazonensis (CHEMBL612877), and "
        "T. cruzi (CHEMBL368) were retrieved from ChEMBL. Compounds containing a "
        "sulfonamide moiety (SMARTS S(=O)(=O)N) were retained. Because these are "
        "whole-cell phenotypic assays in which IC50 and EC50 denote the same "
        "50%-effect quantity, the two endpoints were pooled for classification while "
        "retaining the endpoint label for auditability.")
    h2(doc, "2.2 Curation")
    body(doc,
        "Structures were standardized (largest-fragment selection, charge "
        "neutralization, normalization, and tautomer canonicalization). Potency was "
        "expressed as pIC50, preferring the ChEMBL pchembl_value where available. "
        "Replicate measurements for a compound within an organism were merged to the "
        "median when concordant (<1 log-unit spread) and discarded when discordant. "
        "Compounds were classified Active at pIC50 >= 5.0 (IC50 <= 10 uM). The data "
        "were partitioned by organism into four independent datasets.")
    h2(doc, "2.3 Descriptors and feature selection")
    body(doc,
        "Approximately 1,600 Mordred 2D descriptors were computed and cleaned "
        "(removing high-missingness and constant columns) to 1,271. Continuous "
        "descriptors were preferred over sparse fingerprints for interpretability "
        "and to limit overfitting at small sample size. Per organism, a label-free "
        "variance and correlation filter produced a candidate pool, from which a "
        "stability-selection procedure over cross-validation training folds "
        "(mutual information and random-forest importance) selected 12 descriptors.")
    h2(doc, "2.4 Leakage-free evaluation")
    body(doc,
        "All models were evaluated with repeated (5x5), scaffold-grouped, activity-"
        "stratified cross-validation. Bemis-Murcko scaffolds defined the groups so "
        "that no scaffold - and, since each dataset has one row per molecule, no "
        "molecule - spanned the training and test folds. Feature scaling and "
        "supervised feature selection were refit inside each training fold, so no "
        "information from the held-out fold influenced the pipeline. For reference, "
        "the earlier random split leaked 28.9% of test molecules into training.")
    h2(doc, "2.5 Model panel and final selection")
    body(doc,
        "The panel comprised logistic regression, SVM (RBF), random forest, "
        "gradient boosting, XGBoost, LightGBM, a soft-voting consensus, and a "
        "logistic-regression-meta stacking ensemble, all with class weighting for "
        "imbalance. Hyperparameters were tuned by nested cross-validation. Where the "
        "panel tied within one standard deviation, a parsimony rule selected the "
        "simplest interpretable model to support transparent SHAP interpretation.")
    h2(doc, "2.6 Interpretation and validation")
    body(doc,
        "Model-agnostic SHAP values explained each final model over its 12 "
        "descriptors. Applicability domain was assessed by leverage (Williams plot; "
        "warning leverage h*=3(p+1)/n), and robustness by Y-randomization (50 label "
        "permutations). Analyses used Python 3, RDKit, Mordred, scikit-learn, "
        "XGBoost, LightGBM, and SHAP.")

    # ── Results ──
    h1(doc, "3. Results and Discussion")
    h2(doc, "3.1 Dataset characteristics")
    body(doc,
        "Sulfonamide filtering yielded 690 unique molecules distributed across the "
        "four organisms. The primary L. infantum dataset (n=142) was perfectly "
        "class-balanced (71 active / 71 inactive) and almost entirely IC50-based. "
        "T. cruzi (n=398) was the largest and most active-skewed (70%), while "
        "L. amazonensis (n=72) was the smallest.")
    h2(doc, "3.2 Model performance")
    body(doc,
        "Table 1 reports the honest cross-validated performance of each organism's "
        "final model. Metrics are modest but reflect prospective, leakage-free "
        "estimates on small datasets. The primary L. infantum logistic-regression "
        f"model reached MCC {linf['MCC']} and ROC-AUC {linf['ROC_AUC']}.")
    body(doc, "Table 1. Final per-organism models and leakage-free performance "
              "(mean +/- SD over repeated scaffold-grouped CV).", size=9, italic=True,
         space=2)
    results_table(doc, mt)
    figure(doc, FIG / "panel_performance_summary.png",
           "Figure 1. Leakage-free cross-validated MCC (mean +/- SD) for each "
           "per-organism sulfonamide model; the final model family and dataset size "
           "are annotated.")
    h2(doc, "3.3 No advantage from ensembles or gradient boosting")
    body(doc,
        "Across all four organisms, XGBoost and LightGBM performed comparably to "
        "random forest and SVM, and the consensus/stacking ensembles did not exceed "
        "the best single interpretable model within cross-validation uncertainty. "
        "This supports a parsimonious modeling strategy at this data scale and "
        "yields cleaner, more defensible interpretation.")
    h2(doc, "3.4 Validation")
    body(doc,
        "Applicability-domain coverage ranged from 95.8% to 98.6% of compounds "
        "(Figure 2). In Y-randomization, real models substantially exceeded the "
        "permuted-label distribution (p=0.02 for every organism), confirming that "
        "the models capture genuine signal rather than chance correlations.")
    figure(doc, FIG / "L_infantum_williams.png",
           "Figure 2. Williams plot for the L. infantum model: leverage vs "
           "standardized residuals; nearly all compounds fall within the "
           "applicability domain.", width=5.2)
    h2(doc, "3.5 Structure-activity relationships (L. infantum)")
    body(doc,
        "SHAP analysis of the primary model highlighted atomic-number-weighted "
        f"autocorrelation and partial-charge descriptors ({top_sar}) as the leading "
        "drivers (Figure 3). Grounding these against the data, active compounds tend "
        "to be larger and heavier (higher Z-weighted autocorrelation) and less "
        "dominated by specific polar surface area, whereas halogen (Cl/Br) content "
        "did not distinguish actives from inactives. The relationships are trends "
        "that can guide design rather than deterministic rules.")
    figure(doc, FIG / "L_infantum_shap_beeswarm.png",
           "Figure 3. SHAP summary (beeswarm) for the L. infantum model over its 12 "
           "descriptors; each point is a compound, colored by descriptor value.",
           width=5.4)
    h2(doc, "3.6 Cross-species structure-activity relationships")
    body(doc,
        "Comparing SHAP importances across organisms showed that the descriptor "
        "drivers are largely organism-specific: only one descriptor (PEOE_VSA9) "
        "ranked among the top contributors for more than one organism (L. infantum "
        "and T. cruzi). This heterogeneity is consistent with distinct parasite "
        "biology and provides direct justification for per-organism modeling over a "
        "single pooled classifier.")
    h2(doc, "3.7 Limitations")
    body(doc,
        "The datasets are small (72-398 compounds), so confidence intervals on the "
        "metrics are wide and the abstract descriptors offer interpretive rather "
        "than mechanistic explanations. Activity is a binary threshold on pooled "
        "IC50/EC50 phenotypic data. These honest constraints should temper any "
        "prospective claims.")

    # ── Conclusions ──
    h1(doc, "4. Conclusions and Future Work")
    body(doc,
        "We present a reproducible, leakage-free QSAR panel for anti-kinetoplastid "
        "sulfonamides with L. infantum as the primary target. Under an honest "
        "scaffold-grouped protocol, simple interpretable models matched a full "
        "algorithm panel, achieving modest but trustworthy performance with strong "
        "applicability-domain coverage and significant Y-randomization. SHAP "
        "analysis indicates organism-specific structure-activity relationships. "
        "Future work will apply these models to prospective virtual screening of "
        "commercial libraries with applicability-domain filtering, followed by "
        "experimental validation of prioritized sulfonamides.")

    h1(doc, "References")
    for i, ref in enumerate([
        "World Health Organization. Leishmaniasis fact sheet. WHO, Geneva.",
        "Gaulton, A. et al. The ChEMBL database. Nucleic Acids Res.",
        "Moriwaki, H. et al. Mordred: a molecular descriptor calculator. J. Cheminform.",
        "Bemis, G. W.; Murcko, M. A. The properties of known drugs. 1. Molecular "
        "frameworks. J. Med. Chem.",
        "Pedregosa, F. et al. Scikit-learn: Machine Learning in Python. JMLR.",
        "Lundberg, S. M.; Lee, S.-I. A unified approach to interpreting model "
        "predictions (SHAP). NeurIPS.",
        "Sheridan, R. P. Time-split cross-validation as a method for estimating the "
        "goodness of prospective prediction. J. Chem. Inf. Model.",
        "Landrum, G. A.; Riniker, S. Combining IC50 or Ki values from different "
        "sources is a source of significant noise. J. Chem. Inf. Model. 2024.",
    ], 1):
        body(doc, f"[{i}] {ref}", size=9, space=2)

    doc.save(OUT)
    print(f"Saved manuscript -> {OUT}  ({len(doc.paragraphs)} paragraphs)")


if __name__ == "__main__":
    build()
